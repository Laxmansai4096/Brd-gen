import os
import csv
from typing import Dict, Any

def extract_text_from_file(file_path: str, filename: str) -> Dict[str, Any]:
    ext = os.path.splitext(filename)[1].lower()
    text_content = ""
    summary = ""
    
    try:
        if ext in [".txt", ".md", ".json"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text_content = f.read()
        elif ext == ".csv":
            lines = []
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for i, row in enumerate(reader):
                    if i < 50:  # sample first 50 rows
                        lines.append(", ".join(row))
            text_content = "\n".join(lines)
        elif ext == ".pdf":
            try:
                import fitz  # PyMuPDF
                doc = fitz.open(file_path)
                pages_text = []
                for page_num in range(min(len(doc), 25)):
                    pages_text.append(f"--- Page {page_num + 1} ---\n" + doc[page_num].get_text())
                text_content = "\n".join(pages_text)
            except Exception:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                pages_text = []
                for idx, page in enumerate(reader.pages[:25]):
                    pages_text.append(f"--- Page {idx + 1} ---\n" + (page.extract_text() or ""))
                text_content = "\n".join(pages_text)
        elif ext in [".docx", ".doc"]:
            try:
                import docx
                doc = docx.Document(file_path)
                doc_lines = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        doc_lines.append(para.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            doc_lines.append(" | ".join(row_cells))
                text_content = "\n".join(doc_lines)
            except Exception as doc_err:
                text_content = f"Error reading document: {doc_err}"
        elif ext in [".xlsx", ".xls"]:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                xl_lines = []
                for sheet in wb.worksheets:
                    xl_lines.append(f"--- Sheet: {sheet.title} ---")
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if cells:
                            xl_lines.append(", ".join(cells))
                text_content = "\n".join(xl_lines)
            except Exception as xl_err:
                text_content = f"Error reading spreadsheet: {xl_err}"
        else:
            text_content = f"Unsupported direct text extraction for file format: {ext}"
            
        summary = text_content[:1000] + ("..." if len(text_content) > 1000 else "")
        return {
            "filename": filename,
            "char_count": len(text_content),
            "preview": summary,
            "full_text": text_content[:20000]  # cap for context window
        }
    except Exception as e:
        return {
            "filename": filename,
            "error": str(e),
            "preview": "Failed to parse file.",
            "full_text": ""
        }
